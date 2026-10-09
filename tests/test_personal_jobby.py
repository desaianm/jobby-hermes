"""Synthetic fixtures only; all candidate state is created in temporary directories."""
import pytest
from personal_jobby.core import salary_decision

@pytest.mark.parametrize('raw,origin,currency,unit,kind,low,high,expected',[
 ('CAD 120000–150000','employer','CAD','annual','base',120000,150000,'eligible'),
 ('CAD 100000–150000','employer','CAD','annual','base',100000,150000,'needs_review'),
 ('CAD 90000','employer','CAD','annual','base',90000,90000,'excluded'),
 ('$150000','employer',None,'annual','base',150000,None,'needs_review'),
 ('USD 150000','employer','USD','annual','base',150000,None,'needs_review'),
 ('CAD 150000 estimate','estimate','CAD','annual','base',150000,None,'needs_review'),
 ('up to CAD 150000','employer','CAD','annual','base',None,150000,'needs_review'),
 ('CAD 100 hourly','employer','CAD','hourly','base',100,100,'needs_review'),
 ('CAD 150000 TC','employer','CAD','annual','total',150000,None,'needs_review'),
 ('nan','employer','CAD','annual','base',float('nan'),None,'needs_review'),
 ('inf','employer','CAD','annual','base',float('inf'),None,'needs_review'),
])
def test_salary(raw,origin,currency,unit,kind,low,high,expected):
 assert salary_decision(dict(raw=raw,origin=origin,currency=currency,unit=unit,kind=kind,lower=low,upper=high,verified=True,source_url="https://example.com/jobs/1"),120000)[0]==expected
from personal_jobby.store import Store
from personal_jobby.core import select_facts

def sample():
 return dict(title='AI Engineer',company='Example',job_url='https://example.com/jobs/1',requisition='1',location='Toronto',description='Python RAG',salary=dict(raw='CAD 120000',origin='employer',currency='CAD',unit='annual',kind='base',lower=120000,verified=True))

def test_persistence(tmp_path):
 s=Store(tmp_path); j=s.add(sample()); assert s.add(sample())['id']==j['id']
 with pytest.raises(ValueError): s.transition(j['id'],'applied')
 s.transition(j['id'],'archived'); assert s.add(sample())['id']==j['id']
 s.receipt(j['id'],text='User confirms submitted',timestamp='2026-10-08T12:00:00+00:00')
 assert s.get(j['id'])['status']=='applied'; assert len(s.events(j['id']))==3
 assert (tmp_path.stat().st_mode & 0o777)==0o700
 assert (tmp_path/'workspace.sqlite').stat().st_mode & 0o777==0o600

def test_grounding():
 p={'facts':[dict(id='f1',text='Built Python RAG',source='verified'),dict(id='f2',text='Docker deployment',source='verified')]}
 r=select_facts(p,sample()); assert r['selected_fact_ids'][0]=='f1'; assert set(r['selected_fact_ids'])=={'f1','f2'}
from personal_jobby.documents import prepare
import fitz
import json
from docx import Document

def test_documents(tmp_path):
 s=Store(tmp_path)
 p=dict(name='Test Candidate',contact='Toronto',answers={},facts=[dict(id='e1',section='EDUCATION',group='University | Degree',date='Apr 2024',text='Degree',source='fixture',excerpt='Degree'),dict(id='a1',section='EXPERIENCE',group='Company | AI Engineer',date='Apr 2026 – Present',text='Built Python RAG services.',source='fixture',excerpt='Built Python RAG services.')])
 (tmp_path/'candidate.json').write_text(json.dumps(p)); j=s.add(sample()); meta=prepare(s,j['id']); folder=s.artifact_dir(j['id'])
 d=fitz.open(folder/'resume.pdf'); assert len(d)==1
 assert 'Apr 2026' in d[0].get_text(); assert 'Built Python RAG' in d[0].get_text()
 for block in d[0].get_text('dict')['blocks']:
  for line in block.get('lines',[]):
   for span in line['spans']:
    assert span['size']>=10; assert 'Times' in span['font']; assert span['bbox'][0]>=53; assert span['bbox'][2]<=559
 doc=Document(folder/'resume.docx'); assert doc.sections[0].left_margin.inches==.75
 assert 'Built Python RAG' in '\n'.join(p.text for p in doc.paragraphs)
 assert meta['qa_state']=='pending_visual_review'; assert (folder/'preview.png').exists()
from fastapi.testclient import TestClient
from personal_jobby.app import create_app

def test_security(tmp_path):
 c=TestClient(create_app(tmp_path),base_url='http://localhost')
 assert c.get('/',headers={'host':'evil.example'}).status_code==400
 assert c.post('/api/jobs',json=sample(),headers={'origin':'https://evil.example'}).status_code==403
 assert c.post('/api/jobs',content='{}',headers={'content-type':'text/plain'}).status_code==415
 r=c.get('/'); assert "script-src 'self'" in r.headers['content-security-policy']; assert r.headers['cache-control']=='no-store'
 j=c.post('/api/jobs',json=sample()).json()
 assert c.get(f"/api/jobs/{j['id']}/artifacts/candidate.json").status_code==404
 assert c.get('/api/jobs/1/artifacts/%2E%2E%2Fcandidate.json').status_code==404
 assert c.patch('/api/jobs/1/status',json={'status':'applied'}).status_code==422
 assert c.get('/api/state').json()['submission_adapter']=='unavailable'
from personal_jobby.core import assess, DEFAULTS, Job
from copy import deepcopy

def verified_job():
 j=sample(); j.update(location='Toronto, ON',arrangement='hybrid',arrangement_evidence='Hybrid in Toronto',description='Hybrid in Toronto. Python RAG. Permanent full-time.',employment_type='permanent_full_time',employment_evidence='Permanent full-time')
 j['salary'].update(source_url='https://example.com/jobs/1',raw='CAD 120000 annual base salary')
 return j

def confirmed_settings():
 return {**deepcopy(DEFAULTS),'locations':['Toronto/GTA hybrid','Canada remote'],'criteria_confirmed':True}

@pytest.mark.parametrize('location,arrangement,evidence,expected',[
 ('Toronto, ON','onsite','Onsite in Toronto','needs_review'),
 ('Toronto, ON','unknown','','needs_review'),
 ('Mississauga, ON','hybrid','Hybrid in Mississauga','eligible'),
 ('Markham, ON','hybrid','Hybrid in Markham','eligible'),
 ('Vancouver, Canada','remote','Remote within Canada','eligible'),
 ('Toronto, ON','hybrid','Invented arrangement','needs_review'),
 ('London, ON','hybrid','Hybrid in London','needs_review'),
 ('New York, USA','remote','Remote USA only','needs_review'),
])
def test_arrangement(location,arrangement,evidence,expected):
 j=verified_job(); j.update(location=location,arrangement=arrangement,arrangement_evidence=evidence,description=(evidence if evidence!='Invented arrangement' else 'Python RAG')+' Permanent full-time')
 assert assess(j,confirmed_settings())[0]==expected

def test_excluded_precedence():
 j=verified_job(); j.update(title='Sales Executive',location='',arrangement='unknown');j['salary'].update(lower=90000,upper=90000,raw='CAD 90000 annual base salary')
 assert assess(j,confirmed_settings())[0]=='excluded'

@pytest.mark.parametrize('change',[{'source_url':''},{'verified':False},{'origin':'aggregator'},{'raw':'$120000 per year'},{'raw':'CAD 120000 total compensation'},{'raw':'CAD 150000 annual base salary'}])
def test_salary_provenance(change):
 j=verified_job();j['salary'].update(change)
 assert assess(j,confirmed_settings())[0]=='needs_review'

def test_required_qualifications(tmp_path):
 s=Store(tmp_path);s.save_settings(confirmed_settings())
 p=dict(name='Fixture',contact='',answers={},facts=[dict(id='f1',section='EXPERIENCE',group='Company | AI Engineer',date='Aug 2024 – Present',text='Python RAG',source='fixture',excerpt='Python RAG')])
 (tmp_path/'candidate.json').write_text(json.dumps(p))
 j=verified_job();j['description']+=' Required: 6-10 years of experience in software engineering.'
 result=s.add(j)
 assert result['eligibility']=='needs_review';assert any('years' in r for r in result['explanation'])
 j2=verified_job();j2['job_url']='https://example.com/jobs/2';j2['requisition']='2';j2['description']+=' Required qualifications: PhD in Computer Science.'
 assert s.add(j2)['eligibility']=='needs_review'

def test_resume_preserves_roles_links_and_professional_letter(tmp_path):
 s=Store(tmp_path); j=s.add(verified_job())
 facts=[dict(id='edu',section='EDUCATION',group='University | B.Sc.',date='Apr 2024',text='B.Sc.',source='fixture',excerpt='B.Sc.')]
 for role,date in [('Current | AI Engineer','Apr 2026 – Present'),('Prior | AI Engineer','Aug 2024 – Feb 2026'),('Internship | AI Engineer Intern','Feb 2024 – Jun 2024')]:
  for i in range(8):facts.append(dict(id=role+str(i),section='EXPERIENCE',group=role,date=date,text=('Built Python RAG services with retrieval evaluation, observability, and reliable deployment. ')*2,source='fixture',excerpt='fixture'))
 p=dict(name='Fixture Candidate',contact='Toronto | github.com/example',answers={},links=[{'label':'GitHub','url':'https://github.com/example'},{'label':'LinkedIn','url':'https://linkedin.com/in/example'}],facts=facts)
 (tmp_path/'candidate.json').write_text(json.dumps(p));prepare(s,j['id']);folder=s.artifact_dir(j['id'])
 with fitz.open(folder/'resume.pdf') as pdf:
  text=pdf[0].get_text();assert len(pdf)==1
  for group in ('Current','Prior','Internship'): assert group in text
  assert [text.index(x) for x in ('Current','Prior','Internship')]==sorted(text.index(x) for x in ('Current','Prior','Internship'))
  assert any(l.get('uri')=='https://github.com/example' for l in pdf[0].get_links())
 doc=Document(folder/'resume.docx'); assert any(r.target_ref=='https://github.com/example' for r in doc.part.rels.values())
 with fitz.open(folder/'cover_letter.pdf') as letter:
  text=letter[0].get_text();assert 'Dear Hiring Team' in text;assert 'source-supported' not in text;assert 'Example' in text;assert len(letter)==1
from personal_jobby.import_profile import extract, initialize
from reportlab.pdfgen import canvas

def updated_fixture(path):
 """Generate invented candidate text; no real resume or production metrics."""
 text=['Fixture Candidate','Toronto, ON | fixture@example.com | github.com/example','Experience','Current Co','Toronto','AI Engineer','Apr 2026 - Present','• Reduced median latency from 48.6s to 31.2s in development benchmarks.','• Evaluated three configurations over 24 synthetic test turns.','Prior Co','Toronto','AI Engineer','Aug 2024 - Feb 2026','• Built a synthetic queue service with 42 test messages.','Selected Projects','Queue Demo: Processed synthetic messages with SQLite state.','Technical Skills','Backend & data: Python, FastAPI, PostgreSQL','Cloud & product: GCP, Cloud Run','Education',"University | Bachelor's in Information Technology | Apr 2024"]
 c=canvas.Canvas(str(path));y=780
 for line in text:c.drawString(54,y,line);y-=16
 c.save()

def test_updated_source_migration(tmp_path):
 source=tmp_path/'source.pdf';updated_fixture(source);s=Store(tmp_path/'private')
 old=dict(name='Old',contact='',answers={'phone':'fixture-private-number','work_eligibility':'Canada only'},facts=[{'id':'obsolete','text':'Obsolete AWS metric'}])
 (s.root/'candidate.json').write_text(json.dumps(old))
 j=s.add(sample());folder=s.root/'artifacts'/str(j['id']);folder.mkdir(parents=True);(folder/'metadata.json').write_text(json.dumps({'qa_state':'visual_approved'}))
 result=initialize(s,source,replace=True)
 profile=s.profile();assert profile['answers']==old['answers'];assert 'Obsolete AWS' not in json.dumps(profile['facts'])
 assert len([f for f in profile['facts'] if f['section']=='EXPERIENCE'])==3
 assert 'in development benchmarks' in profile['facts'][0]['text']
 assert len(profile['source_hashes'])==1;assert all(f['source']==str(source) and f['excerpt'] for f in profile['facts'])
 assert result['backup_created'];assert not (folder/'metadata.json').exists()
 assert list((s.root/'backups').glob('candidate-*.json'))[0].stat().st_mode & 0o777==0o600
 assert 'fixture-private-number' not in json.dumps(result)

def test_phone_secrecy(tmp_path):
 c=TestClient(create_app(tmp_path),base_url='http://localhost');s=c.app.state.store
 (s.root/'candidate.json').write_text(json.dumps(dict(name='Fixture',contact='Toronto',answers={'phone':'','work_eligibility':'Canada'},facts=[])))
 phone='+1 416 555 0199'
 r=c.put('/api/profile/phone',json={'phone':phone});assert r.status_code==200;assert phone not in r.text
 state=c.get('/api/state');assert state.json()['phone_present'];assert phone not in state.text;assert 'phone' not in state.json()['profile']['answers']
 c.put('/api/profile/answers',json={'work_eligibility':'Canada only'});assert s.profile()['answers']['phone']==phone
 c.put('/api/profile/phone',json={'phone':''});assert s.profile()['answers']['phone']==phone
 invalid=c.put('/api/profile/phone',json={'phone':phone*20});assert invalid.status_code==422;assert phone not in invalid.text
 c.put('/api/profile/phone',json={'clear':True});assert not c.get('/api/state').json()['phone_present']
 assert 'type="password"' in c.get('/').text;assert 'name="phone"' not in c.get('/').text.split('<form id="answers">')[1].split('</form>')[0]
 assert (s.root/'candidate.json').stat().st_mode & 0o777==0o600

@pytest.mark.parametrize('employment_type,expected',[('permanent_full_time','eligible'),('contract','excluded'),('temporary','excluded'),('part_time','excluded'),('unknown','needs_review')])
def test_permanent_full_time(employment_type,expected):
 j=verified_job();j.update(employment_type=employment_type,employment_evidence='Permanent full-time' if employment_type=='permanent_full_time' else employment_type);j['description']+=' '+j['employment_evidence']
 assert assess(j,confirmed_settings())[0]==expected

def test_current_employer_exclusion():
 j=verified_job();j['company']='cUrReNtCo';settings=confirmed_settings();settings['blacklist']=['CurrentCo']
 assert assess(j,settings)[0]=='excluded'

def test_unknown_requirements_and_fit_update(tmp_path):
 c=TestClient(create_app(tmp_path),base_url='http://localhost');s=c.app.state.store;s.save_settings(confirmed_settings())
 j=verified_job();j['description']+=' Required qualifications: proven enterprise architecture experience.'
 result=c.post('/api/jobs',json=j).json();assert result['eligibility']=='needs_review'
 r=c.put('/api/jobs/'+str(result['id'])+'/fit',json={'arrangement':'hybrid','arrangement_evidence':'Hybrid in Toronto','employment_type':'permanent_full_time','employment_evidence':'Permanent full-time','qualifications_reviewed':True,'qualification_review_note':'Reviewed original requirements against verified facts; no unsupported claims.'})
 assert r.status_code==200;assert r.json()['eligibility']=='eligible'
 assert any(e['kind']=='fit_verification' for e in s.events(result['id']))

def test_search_is_bounded_and_unverified(tmp_path,monkeypatch):
 import pandas as pd
 import jobspy
 from personal_jobby.search import discover
 calls=[]
 def fake(**kwargs):
  calls.append(kwargs);return pd.DataFrame([{'title':'AI Engineer','company':'Fixture','job_url':'https://example.com/fake','location':'Toronto, Canada','description':'Hybrid. Permanent full-time. Python','currency':'CAD','min_amount':150000,'max_amount':200000,'interval':'yearly','salary_source':'direct_data'}])
 monkeypatch.setattr(jobspy,'scrape_jobs',fake);s=Store(tmp_path)
 r=discover(s,'AI Engineer',1);assert len(r['jobs'])==1;assert r['jobs'][0]['eligibility']=='needs_review';assert r['jobs'][0]['salary']['origin']=='aggregator'
 assert calls[0]['country_indeed']=='Canada'
 with pytest.raises(ValueError): discover(s,'AI Engineer',26)

def test_stale_artifacts_and_symlink_boundary(tmp_path):
 c=TestClient(create_app(tmp_path),base_url='http://localhost');s=c.app.state.store
 (s.root/'candidate.json').write_text(json.dumps(dict(name='Fixture',contact='Toronto',answers={},facts=[dict(id='f',section='EXPERIENCE',group='Co | AI Engineer',date='Apr 2026 - Present',text='Built Python services.',source='fixture',excerpt='Built Python services.')],source_hashes={'fixture':'first'})))
 j=s.add(sample());prepare(s,j['id']);folder=s.artifact_dir(j['id'])
 profile=s.profile();profile['source_hashes']={'fixture':'second'};(s.root/'candidate.json').write_text(json.dumps(profile))
 assert c.get('/api/jobs/1/artifacts/resume.pdf').status_code==409
 assert 'Documents source changed' in ' '.join(s.get(j['id'])['blockers'])
 outside=tmp_path/'outside';outside.mkdir();(outside/'resume.pdf').write_bytes(b'private')
 import shutil
 shutil.rmtree(folder);folder.symlink_to(outside,target_is_directory=True)
 assert c.get('/api/jobs/1/artifacts/resume.pdf').status_code==404

def test_phone_never_in_history_or_validation(tmp_path):
 c=TestClient(create_app(tmp_path),base_url='http://localhost');s=c.app.state.store;s.save_phone('+1 416 555 0100')
 j=s.add(sample());s.receipt(j['id'],text='Receipt for +1 416 555 0100',timestamp='2026-10-08T12:00:00+00:00')
 assert '+1 416 555 0100' not in c.get('/api/jobs/1').text
 assert '+1 416 555 0100' not in json.dumps(s.events(1))
 bad=c.put('/api/profile/answers',json={'phone':'+1 416 555 0100'});assert bad.status_code==422;assert '+1 416 555 0100' not in bad.text

def test_dashboard_workflow_controls(tmp_path):
 c=TestClient(create_app(tmp_path),base_url='http://localhost');html=c.get('/').text;js=c.get('/static/app.js').text
 for identifier in ('fit','phone','salary','receipt','approval'): assert 'id="'+identifier+'"' in html
 assert 'name="employment_type"' in html;assert 'name="arrangement"' in html
 assert 'onclick=' not in html;assert 'innerHTML' not in js
 assert 'qualifications_reviewed' in html

def test_cli_help_and_status_no_phone(tmp_path,monkeypatch,capsys):
 import personal_jobby.__main__ as cli
 monkeypatch.setattr(cli,'Store',lambda:Store(tmp_path));monkeypatch.setattr('sys.argv',['personal_jobby','status'])
 s=Store(tmp_path);s.save_phone('+1 416 555 0111');cli.main();out=capsys.readouterr().out
 assert '+1 416 555 0111' not in out;assert 'phone_present' in out;assert 'unavailable' in out

@pytest.mark.parametrize('raw,lower,upper,expected',[
 ('CAD 100000–150000 annual base salary',150000,None,'needs_review'),
 ('CAD 120000–150000 annual base salary',150000,150000,'needs_review'),
 ('CAD 120–150K annual base salary',120000,150000,'eligible'),
 ('Up to CAD 150000 annual base salary',150000,None,'needs_review'),
 ('CAD 120000 annual base salary plus CAD 20000 bonus',120000,None,'needs_review'),
])
def test_raw_salary_bound_mapping(raw,lower,upper,expected):
 s=verified_job()['salary'];s.update(raw=raw,lower=lower,upper=upper)
 assert salary_decision(s,120000)[0]==expected

def test_current_employer_independent_of_blacklist():
 j=verified_job();j['company']='CURRENT CO Inc.'
 p={'facts':[dict(section='EXPERIENCE',group='Current Co | AI Engineer',employer='Current Co',date='Apr 2026 - Present',text='Python')]}
 settings=confirmed_settings();settings['blacklist']=[]
 assert assess(j,settings,p)[0]=='excluded'
 j['company']='Unrelated Current Co Research'; assert assess(j,settings,p)[0]=='eligible'

@pytest.mark.parametrize('text',['Not permanent full-time','Permanent full-time. This is a six-month contract role.','Permanent full-time. Temporary employment.'])
def test_negated_or_conflicting_employment(text):
 j=verified_job();j['description']='Hybrid in Toronto. '+text;j['employment_evidence']=text
 assert assess(j,confirmed_settings())[0]!='eligible'

def test_skill_years_and_qualification_heading():
 j=verified_job();j['description']+=' Qualifications: 5 years of Rust experience. Python preferred. Required: PhD.'
 p={'facts':[dict(section='EXPERIENCE',group='Other Co',date='Jan 2010 - Present',text='Python backend services.')]}
 state,reasons=assess(j,confirmed_settings(),p);assert state=='needs_review';assert any('skill-specific' in r for r in reasons);assert any('PhD' in r for r in reasons)
 j['qualifications_reviewed']=True;j['qualification_review_note']='Reviewed';assert assess(j,confirmed_settings(),p)[0]=='needs_review'

def test_phone_format_variants_all_paths(tmp_path):
 c=TestClient(create_app(tmp_path),base_url='http://localhost');s=c.app.state.store
 p=dict(name='Fixture',contact='Toronto',answers={},facts=[dict(id='f',section='EXPERIENCE',group='Co | Engineer',date='Apr 2026 - Present',text='Built Python services.',source='fixture',excerpt='Built Python services.')])
 (s.root/'candidate.json').write_text(json.dumps(p));s.save_phone('+1 416 555 0199')
 j=sample();j['description']='Call 416-555-0199 or 14165550199.';created=c.post('/api/jobs',json=j).json()
 assert '416-555-0199' not in json.dumps(created);assert '14165550199' not in json.dumps(created)
 s.save_phone('+1 416 555 0101');s.save_phone(clear=True)
 status=c.patch('/api/jobs/1/status',json={'status':'archived'});assert '416-555-0199' not in status.text
 prepare(s,1);posting=c.get('/api/jobs/1/artifacts/posting.json');assert posting.status_code==200;assert '416-555-0199' not in posting.text;assert '14165550199' not in posting.text
 assert '416-555-0199' not in c.get('/api/state').text
 bad=c.put('/api/jobs/1/salary',json={'raw':'416-555-0199','lower':'416-555-0199'});assert bad.status_code==422;assert '416-555-0199' not in bad.text
 assert '416-555-0199' not in json.dumps(s.events(1))

def test_private_directory_root_read_denied(tmp_path,monkeypatch):
 import errno,os,stat
 from personal_jobby import private_fs
 if not hasattr(os,'O_PATH'): pytest.skip('O_PATH unavailable')
 real_open=os.open;root_flags=[]
 def restricted_open(path,flags,*args,**kwargs):
  if path=='/':
   root_flags.append(flags)
   if not flags & os.O_PATH: raise PermissionError(errno.EACCES,'Permission denied',path)
  return real_open(path,flags,*args,**kwargs)
 monkeypatch.setattr(os,'open',restricted_open)
 root=tmp_path/'private';private_fs.mkdir(root)
 private_fs.write_text(root/'fixture.txt','fixture')
 assert private_fs.read_text(root/'fixture.txt')=='fixture'
 assert stat.S_IMODE(root.stat().st_mode)==0o700
 assert root_flags and all(flags & os.O_PATH for flags in root_flags)

def test_private_directory_without_o_path(tmp_path,monkeypatch):
 import os
 from personal_jobby import private_fs
 monkeypatch.delattr(os,'O_PATH',raising=False)
 root=tmp_path/'private';private_fs.mkdir(root)
 private_fs.write_text(root/'fixture.txt','fixture')
 assert private_fs.read_text(root/'fixture.txt')=='fixture'

def test_private_directory_root_descriptor_readable():
 import os
 from personal_jobby import private_fs
 fd=private_fs.directory('/')
 try: os.fsync(fd)
 finally: os.close(fd)

def test_private_reads_writes_reject_symlinks(tmp_path):
 from personal_jobby.store import private_write
 outside=tmp_path/'outside';outside.mkdir();victim=outside/'victim.json';victim.write_text('{"secret":"untouched"}')
 root=tmp_path/'private';s=Store(root);(root/'candidate.json').symlink_to(victim)
 with pytest.raises(ValueError): s.profile()
 with pytest.raises(ValueError): private_write(root/'candidate.json','changed')
 assert victim.read_text()=='{"secret":"untouched"}'
 (root/'candidate.json').unlink();(root/'backups').symlink_to(outside,target_is_directory=True)
 source=tmp_path/'source.pdf';updated_fixture(source);(root/'candidate.json').write_text(json.dumps({'answers':{},'facts':[]}))
 with pytest.raises(ValueError): initialize(s,source,replace=True)
 linked=tmp_path/'linked';linked.symlink_to(outside,target_is_directory=True)
 with pytest.raises(ValueError): Store(linked)
 dbroot=tmp_path/'dbroot';dbroot.mkdir();(dbroot/'workspace.sqlite').symlink_to(victim)
 with pytest.raises(ValueError): Store(dbroot)

def test_prepare_rejects_linked_artifact_directory(tmp_path):
 s=Store(tmp_path/'private');outside=tmp_path/'outside';outside.mkdir()
 p=dict(name='Fixture',contact='',answers={},facts=[dict(id='f',section='EXPERIENCE',group='Co',date='Apr 2026 - Present',text='Python',source='fixture',excerpt='Python')])
 (s.root/'candidate.json').write_text(json.dumps(p));j=s.add(sample());(s.root/'artifacts').symlink_to(outside,target_is_directory=True)
 with pytest.raises(ValueError): prepare(s,j['id'])
 assert not list(outside.iterdir())

def document_workspace(tmp_path):
 s=Store(tmp_path);p=dict(name='Fixture',contact='Toronto',answers={},facts=[dict(id='f',section='EXPERIENCE',group='Company | AI Engineer',date='Apr 2026 - Present',text='Built Python services.',source='fixture',excerpt='Built Python services.')],source_hashes={'fixture':'v1'})
 (s.root/'candidate.json').write_text(json.dumps(p));j=s.add(verified_job());return s,j

def test_visual_approval_binds_all_versions(tmp_path):
 s,j=document_workspace(tmp_path);c=TestClient(create_app(tmp_path),base_url='http://localhost')
 meta=prepare(s,j['id']);token=meta['review_token']
 approve=lambda t:c.post('/api/jobs/1/visual-approval',json={'reviewer':'External reviewer','note':'Viewed PDF and DOCX','review_token':t})
 assert approve(token).status_code==200
 p=s.profile();p['facts'][0]['text']='Built Python evaluation services.';(s.root/'candidate.json').write_text(json.dumps(p))
 assert approve(token).status_code==409;assert c.get('/api/jobs/1/artifacts/resume.pdf').status_code==409
 meta=prepare(s,1);settings=s.settings();settings['minimum']=130000;s.save_settings(settings)
 assert approve(meta['review_token']).status_code==409
 meta=prepare(s,1);c.put('/api/jobs/1/fit',json={'arrangement':'onsite','arrangement_evidence':'Hybrid in Toronto','employment_type':'unknown'})
 assert approve(meta['review_token']).status_code==409
 meta=prepare(s,1);folder=s.artifact_dir(1);(folder/'resume.pdf').write_bytes(b'changed')
 assert approve(meta['review_token']).status_code==409
 meta=prepare(s,1);assert approve(token).status_code==409
 assert approve(meta['review_token']).status_code==200

def test_failed_generation_invalidates_before_writes(tmp_path,monkeypatch):
 import personal_jobby.documents as docs
 s,j=document_workspace(tmp_path);c=TestClient(create_app(tmp_path),base_url='http://localhost')
 meta=prepare(s,1);assert c.post('/api/jobs/1/visual-approval',json={'reviewer':'Reviewer','note':'Viewed','review_token':meta['review_token']}).status_code==200
 def fail(*args,**kwargs):
  assert c.get('/api/jobs/1/artifacts/resume.pdf').status_code==409
  raise ValueError('Generation failed')
 monkeypatch.setattr(docs,'pdf',fail)
 with pytest.raises(ValueError):prepare(s,1)
 assert c.get('/api/jobs/1/artifacts/resume.pdf').status_code==409
 assert s.get(1)['materials']['qa_state']!='visual_approved'

def test_bounded_keyword_substitutions_grounded_documents(tmp_path):
 s,j=document_workspace(tmp_path)
 p=s.profile();text='Built retrieval-augmented generation (RAG) and Model Context Protocol (MCP) APIs; reduced latency from 48.6s to 31.2s in development benchmarks, with advisory checks and agent review.'
 p['facts'][0]['text']=text;p['facts'][0]['excerpt']=text;(s.root/'candidate.json').write_text(json.dumps(p))
 job=verified_job();job['description']+=' RAG, MCP, application programming interfaces. AWS, PostgreSQL, PyTorch, TensorFlow, GPT, 5 years.'
 selection=select_facts(p,job);tailored=selection['tailoring'][0]
 assert tailored['original']==text;assert 'retrieval-augmented generation' not in tailored['rendered'];assert '(RAG) (RAG)' not in tailored['rendered']
 assert 'application programming interfaces' in tailored['rendered'];assert len(tailored['substitutions'])==3
 for qualifier in ('48.6s','31.2s','in development benchmarks','advisory checks','agent review'):assert qualifier in tailored['rendered']
 for keyword in ('AWS','PostgreSQL','PyTorch','TensorFlow','GPT','5 years'):assert keyword not in tailored['rendered']
 no_match=verified_job();no_match['description']='Hybrid in Toronto. Permanent full-time.';assert select_facts(p,no_match)['tailoring'][0]['substitutions']==[]
 # The posting changed only in a temporary fixture database.
 with s.connect() as connection:connection.execute('UPDATE jobs SET data=? WHERE id=1',(json.dumps(Job.model_validate(job).model_dump()),))
 meta=prepare(s,1);folder=s.artifact_dir(1)
 with fitz.open(folder/'resume.pdf') as pdf:assert 'application programming interfaces' in pdf[0].get_text()
 evidence=json.loads((folder/'evidence.json').read_text());assert evidence['tailoring'][0]['original']==text
 assert meta['selection_mode']=='source facts with bounded equivalent wording'

def test_project_links_associated_by_geometry(tmp_path):
 path=tmp_path/'projects.pdf';c=canvas.Canvas(str(path));y=780;positions={}
 rows=['Fixture','Toronto | fixture@example.com','Experience','Company','Toronto','Engineer','Apr 2026 - Present','• Built Python.','Selected Projects','Alpha: Built a queue. GitHub','Beta: Built a voice tool. GitHub','Technical Skills','Backend: Python','Education','University | Degree | Apr 2024']
 for row in rows:
  c.drawString(54,y,row);positions[row]=y;y-=20
 for title,url in [('Beta: Built a voice tool. GitHub','https://github.com/example/beta'),('Alpha: Built a queue. GitHub','https://github.com/example/alpha')]:
  ypos=positions[title];c.linkURL(url,(54,ypos-2,310,ypos+12))
 c.save();p=extract(path);projects=[f for f in p['facts'] if f['section']=='PROJECTS']
 assert [f['link'] for f in projects]==['https://github.com/example/alpha','https://github.com/example/beta']

def test_canonical_alias_deduplicates_real_listing(tmp_path):
 s=Store(tmp_path);j=sample();j['original_source']={'job_url_direct':'https://careers.example.com/jobs/abc'};first=s.add(j)
 direct=sample();direct['job_url']='https://careers.example.com/jobs/abc';direct['requisition']='';assert s.add(direct)['id']==first['id'];assert len(s.jobs())==1

def test_inline_skills_and_first_person_cover(tmp_path):
 s,j=document_workspace(tmp_path);p=s.profile();p['facts'].append(dict(id='skill',section='TECHNICAL SKILLS',group='LLM engineering',date='',text='Prompt engineering, tool calling, structured outputs, retrieval-augmented generation (RAG), multimodal extraction, evaluation',source='fixture',excerpt='fixture'));(s.root/'candidate.json').write_text(json.dumps(p));prepare(s,1)
 folder=s.artifact_dir(1)
 with fitz.open(folder/'resume.pdf') as d:
  text=d[0].get_text();assert 'LLM engineering:' in text;assert '\nevaluation\n' not in text
 with fitz.open(folder/'cover_letter.pdf') as d:assert 'I built Python services.' in d[0].get_text()

@pytest.mark.parametrize('ending',['Pydantic','Next.js'])
def test_inline_skills_keep_final_tokens_together(tmp_path,ending):
 from personal_jobby.documents import pdf, docx
 body=', '.join(['Python']*11+['FastAPI',ending])
 content=[('skill',('Backend & data',body))]
 pdf(tmp_path/'skills.pdf',content)
 with fitz.open(tmp_path/'skills.pdf') as d:
  assert len(d)==1;assert tuple(d[0].rect)==(0,0,612,792)
  text=d[0].get_text()
  assert text.splitlines()[-1].split()==['FastAPI,',ending]
  assert ' '.join(text.split())=='Backend & data: '+body
  for block in d[0].get_text('dict')['blocks']:
   for line in block.get('lines',[]):
    for span in line['spans']:
     assert span['size']>=10;assert span['bbox'][0]>=54;assert span['bbox'][2]<=558
 docx(tmp_path/'skills.docx',content)
 doc=Document(tmp_path/'skills.docx')
 assert doc.paragraphs[0].text=='Backend & data: '+body.rsplit(' ',1)[0]+'\u00a0'+ending
 assert doc.styles['Normal'].font.size.pt==10
 sec=doc.sections[0]
 assert (sec.page_width.inches,sec.page_height.inches)==(8.5,11)
 assert all(getattr(sec,side+'_margin').inches==.75 for side in ('top','bottom','left','right'))

def test_cli_brief_omits_descriptions(tmp_path,monkeypatch,capsys):
 import personal_jobby.__main__ as cli
 s=Store(tmp_path);s.add(sample());monkeypatch.setattr(cli,'Store',lambda:s);monkeypatch.setattr('sys.argv',['personal_jobby','status','--brief']);cli.main();out=capsys.readouterr().out
 assert 'description' not in out;assert 'job_count' in out

def test_application_manifest_and_historical_download(tmp_path):
 from personal_jobby.audit import create_attempt,add_screenshot
 s,j=document_workspace(tmp_path);meta=prepare(s,1)
 c=TestClient(create_app(tmp_path),base_url='http://localhost')
 c.post('/api/jobs/1/visual-approval',json={'reviewer':'Reviewer','note':'Viewed both documents','review_token':meta['review_token']})
 manifest=dict(operator='hermes',generation_id=meta['generation_id'],review_token=meta['review_token'],application_url='https://example.com/apply',resume={'filename':'resume.pdf','sha256':meta['artifact_hashes']['resume.pdf']},cover_letter={'filename':'cover_letter.docx','sha256':meta['artifact_hashes']['cover_letter.docx']})
 attempt=create_attempt(s,1,manifest);assert s.get(1)['status']!='applied';assert attempt['cover_letter_used'] is True
 from PIL import Image
 image=tmp_path/'screen.png';Image.new('RGB',(100,100),'white').save(image)
 with pytest.raises(ValueError):add_screenshot(s,attempt['id'],image,'filled_form','Reviewed',False)
 pre=add_screenshot(s,attempt['id'],image,'filled_form','Checked filled fields with phone masked before capture',True)
 add_screenshot(s,attempt['id'],image,'employer_confirmation','Checked employer confirmation; phone omitted',True)
 result=s.receipt(1,text='Employer confirmed submission',timestamp='2026-10-08T12:00:00-04:00',attempt_id=attempt['id']);assert result['submitted_at']=='2026-10-08T12:00:00-04:00'
 prepare(s,1);settings=s.settings();settings['minimum']=140000;s.save_settings(settings)
 history=c.get('/api/jobs/1').json()['job']['attempts'];assert history[0]['generation_id']==meta['generation_id'];assert history[0]['status']=='applied'
 response=c.get('/api/jobs/1/attempts/'+attempt['id']+'/documents/resume.pdf');assert response.status_code==200
 import hashlib
 assert hashlib.sha256(response.content).hexdigest()==manifest['resume']['sha256']
 assert c.get('/api/jobs/1/attempts/'+attempt['id']+'/screenshots/'+pre['id']).status_code==200
 assert (s.root/'applications'/attempt['id']/'documents'/'resume.pdf').stat().st_mode & 0o777==0o600
 with pytest.raises(ValueError): create_attempt(s,1,manifest)

def test_resume_only_manifest_requires_explicit_cover_decision(tmp_path):
 from personal_jobby.audit import create_attempt
 s,j=document_workspace(tmp_path);meta=prepare(s,1)
 c=TestClient(create_app(tmp_path),base_url='http://localhost')
 c.post('/api/jobs/1/visual-approval',json={'reviewer':'Reviewer','note':'Viewed fixture documents','review_token':meta['review_token']})
 manifest=dict(operator='hermes',generation_id=meta['generation_id'],review_token=meta['review_token'],resume={'filename':'resume.pdf','sha256':meta['artifact_hashes']['resume.pdf']},cover_letter=None)
 attempt=create_attempt(s,1,manifest)
 assert attempt['cover_letter_used'] is False
 assert attempt['documents']=={'resume':manifest['resume']}
 assert attempt['documents_used']=='explicit manifest (cover letter: none)'
 folder=s.root/'applications'/attempt['id']
 assert {p.name for p in (folder/'documents').iterdir()}=={'resume.pdf'}
 assert json.loads((folder/'manifest.json').read_text())['cover_letter_used'] is False
 history=c.get('/api/jobs/1').json()['job']['attempts'][0]
 assert history['documents']==attempt['documents'];assert history['cover_letter_used'] is False
 assert c.get('/api/jobs/1/attempts/'+attempt['id']+'/documents/resume.pdf').status_code==200
 assert c.get('/api/jobs/1/attempts/'+attempt['id']+'/documents/cover_letter.pdf').status_code==404
 del manifest['cover_letter']
 with pytest.raises(ValueError,match='Explicit cover_letter'):create_attempt(s,1,manifest)
 manifest['cover_letter']=None;manifest['resume']=None
 with pytest.raises(ValueError):create_attempt(s,1,manifest)

def test_manual_receipt_documents_unknown(tmp_path):
 s,j=document_workspace(tmp_path);prepare(s,1);s.receipt(1,text='User reports confirmed',timestamp='2026-10-08T12:00:00+00:00')
 attempt=s.get(1)['attempts'][0];assert attempt['documents_used']=='unknown';assert attempt['generation_id'] is None;assert attempt['documents']=={}
 assert attempt['cover_letter_used'] is None


def test_managed_attempt_rejects_wrong_hash_and_unconfirmed_receipt(tmp_path):
 from personal_jobby.audit import create_attempt
 s,j=document_workspace(tmp_path);meta=prepare(s,1);c=TestClient(create_app(tmp_path),base_url='http://localhost')
 manifest=dict(operator='hermes',generation_id=meta['generation_id'],review_token=meta['review_token'],application_url='https://example.com/apply',resume={'filename':'resume.pdf','sha256':'0'*64},cover_letter={'filename':'cover_letter.pdf','sha256':meta['artifact_hashes']['cover_letter.pdf']})
 with pytest.raises(ValueError):create_attempt(s,1,manifest)
 c.post('/api/jobs/1/visual-approval',json={'reviewer':'Reviewer','note':'Viewed','review_token':meta['review_token']})
 with pytest.raises(ValueError):create_attempt(s,1,manifest)
 manifest['resume']['sha256']=meta['artifact_hashes']['resume.pdf'];attempt=create_attempt(s,1,manifest)
 with pytest.raises(ValueError):s.receipt(1,text='Confirmation',timestamp='2026-10-08T12:00:00+00:00',attempt_id=attempt['id'])
 assert s.get(1)['status']!='applied'

def test_audit_cli_and_ui_are_explicit(tmp_path,monkeypatch,capsys):
 import personal_jobby.__main__ as cli
 monkeypatch.setattr('sys.argv',['personal_jobby','screenshot','--help'])
 with pytest.raises(SystemExit) as exit:cli.main()
 assert exit.value.code==0;help=capsys.readouterr().out;assert '--phone-redacted' in help;assert '--qa-note' in help
 c=TestClient(create_app(tmp_path),base_url='http://localhost');html=c.get('/').text;js=c.get('/static/app.js').text
 assert 'id="attempts"' in html;assert 'Submitted at:' in js;assert 'documents_used' in js

def test_screenshot_rejects_invalid_or_oversized_input(tmp_path):
 from personal_jobby.audit import create_attempt,add_screenshot
 s,j=document_workspace(tmp_path);attempt=create_attempt(s,1,{'operator':'user_reported'})
 bad=tmp_path/'bad.png';bad.write_bytes(b'not an image')
 with pytest.raises(ValueError):add_screenshot(s,attempt['id'],bad,'filled_form','Reviewed',True)
 with bad.open('wb') as f:f.truncate(8*1024*1024+1)
 with pytest.raises(ValueError):add_screenshot(s,attempt['id'],bad,'filled_form','Reviewed',True)

def test_changed_authoritative_source_requires_reimport(tmp_path):
 s=Store(tmp_path/'private');source=tmp_path/'source.pdf';updated_fixture(source);initialize(s,source);s.add(verified_job());prepare(s,1)
 source.write_bytes(source.read_bytes()+b'changed')
 with pytest.raises(ValueError): prepare(s,1)
 assert s.get(1)['materials']['qa_state']!='visual_approved'

def test_toronto_borough_is_gta_hybrid():
 j=verified_job();j['location']='North York, ON, Canada'
 assert assess(j,confirmed_settings())[0]=='eligible'

def test_dashboard_narrow_viewport_css_contract(tmp_path):
 import re
 c=TestClient(create_app(tmp_path),base_url='http://localhost')
 css=c.get('/static/style.css').text
 def declarations(selector, stylesheet=css):
  match=re.search(re.escape(selector)+r'\{([^{}]*)\}',stylesheet)
  assert match is not None
  return dict(part.strip().split(':',1) for part in match[1].split(';') if part.strip())
 # Inherited wrapping covers identity text, attempt generation IDs and hashes.
 assert declarations(':root').get('overflow-wrap')=='anywhere'
 assert declarations('*').get('box-sizing')=='border-box'
 controls=declarations('input,textarea,select,button')
 assert controls.get('max-width')=='100%'
 assert controls.get('min-width')=='0'
 mobile=css.split('@media(max-width:600px)',1)[1]
 buttons=declarations('button',mobile)
 assert buttons.get('width')=='100%'
 # A full-width button plus its old 8px right margin exceeds its container.
 assert buttons.get('margin-right')=='0'
 assert not re.search(r'overflow(?:-x)?\s*:\s*(?:hidden|clip)',css)
