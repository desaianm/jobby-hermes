"""Grounded one-page documents. Only complete source facts may be selected."""
import hashlib
import json
import os
from collections import OrderedDict, Counter
from xml.sax.saxutils import escape
import fitz
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from personal_jobby.core import select_facts, public_url
from personal_jobby.store import private_write, now
from personal_jobby import private_fs as fs

NAMES={'resume.pdf','resume.docx','cover_letter.pdf','cover_letter.docx','evidence.json','metadata.json','posting.json','preview.png','cover_preview.png'}


def contact_parts(profile):
    links=profile.get('links',[]);parts=[]
    for part in profile['contact'].split('|'):
        part=part.strip()
        if not any(l['url'].removeprefix('https://').removeprefix('http://') in part or l['url'] in part for l in links): parts.append((part,''))
    for link in links:
        public_url(link['url']);parts.append((link['label'],link['url']))
    return parts


def lines(profile,facts):
    out=[('name',profile['name']),('contact',contact_parts(profile))]
    for section in ('EDUCATION','EXPERIENCE','PROJECTS','TECHNICAL SKILLS'):
        fs=[f for f in facts if f['section']==section]
        if not fs: continue
        out.append(('section',section));groups=OrderedDict()
        for f in profile['facts']:
            if f['section']==section and any(x['id']==f['id'] for x in fs): groups.setdefault((f['group'],f['date']),[])
        for f in fs: groups[(f['group'],f['date'])].append(f)
        for (group,date),items in groups.items():
            if section=='TECHNICAL SKILLS':
                for f in items: out.append(('skill',(group,f['text'])))
                continue
            out.append(('heading',(group,date)))
            for f in items:
                out.append(('body',f['text'] if section in ('EDUCATION','TECHNICAL SKILLS') else '• '+f['text']))
                if f.get('link'): out.append(('contact',[('GitHub',f['link'])]))
    return out


def keep_skill_tail(body):
    """Keep the final two skill tokens together without changing source words."""
    head,space,tail=body.rpartition(' ')
    return head+'\u00a0'+tail if space else body


def pdf(path,content,leading=12.5,body_space=3):
    styles={k:ParagraphStyle(k,fontName='Times-Bold' if k in ('name','section','heading') else 'Times-Roman',fontSize=14 if k=='name' else 10,leading=leading,spaceAfter=4 if k!='body' else body_space) for k in ('name','section','heading','body','contact','skill')}
    styles['heading'].keepWithNext=True;styles['section'].keepWithNext=True
    story=[]
    for kind,text in content:
        if kind=='section': story.extend([Spacer(1,6),HRFlowable(width='100%',thickness=.5,color='black')])
        if kind=='heading':
            group,date=text
            if date:
                right=ParagraphStyle('right',parent=styles[kind],alignment=2)
                table=Table([[Paragraph(escape(group),styles[kind]),Paragraph(escape(date),right)]],colWidths=[365,127])
                table.setStyle(TableStyle([('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0),('VALIGN',(0,0),(-1,-1),'TOP')]))
                table.keepWithNext=True;story.append(table);continue
            text=escape(group)
        elif kind=='skill':
            group,body=text;text='<b>'+escape(group)+':</b> '+escape(keep_skill_tail(body))
        elif kind=='contact':
            text=' | '.join('<a href="'+escape(url,{'"':'&quot;'})+'" color="black">'+escape(label)+'</a>' if url else escape(label) for label,url in text)
        else: text=escape(text)
        story.append(Paragraph(text,styles[kind]))
    SimpleDocTemplate(str(path),pagesize=(612,792),leftMargin=54,rightMargin=54,topMargin=54,bottomMargin=54).build(story)
    fs.chmod(path)


def hyperlink(p,label,url):
    public_url(url)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');props.append(font)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');props.append(size);run.append(props)
    text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);p._p.append(link)


def docx(path,content,leading=12.5,body_space=3):
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(body_space);normal.paragraph_format.line_spacing=Pt(leading)
    for kind,text in content:
        p=doc.add_paragraph()
        if kind=='skill':
            group,body=text;p.add_run(group+': ').bold=True;p.add_run(keep_skill_tail(body))
        elif kind=='contact':
            for i,(label,url) in enumerate(text):
                if i: p.add_run(' | ')
                if url: hyperlink(p,label,url)
                else: p.add_run(label)
        else:
            if kind=='heading':
                group,date=text;text=group+('\t'+date if date else '')
                p.paragraph_format.tab_stops.add_tab_stop(Inches(6.83),WD_TAB_ALIGNMENT.RIGHT)
                p.paragraph_format.keep_with_next=True
            r=p.add_run(text);r.bold=kind!='body';r.font.size=Pt(14 if kind=='name' else 10)
        if kind=='section':
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
            borders=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'4');borders.append(bottom);p._p.get_or_add_pPr().append(borders)
    doc.save(path);fs.chmod(path)


def qa(path,preview):
    with fitz.open(path) as d:
        if len(d)!=1: raise ValueError('Document exceeds one page')
        if not d[0].get_text().strip(): raise ValueError('Missing ATS text')
        for b in d[0].get_text('dict')['blocks']:
            for l in b.get('lines',[]):
                for sp in l['spans']:
                    if sp['size']<10 or 'Times' not in sp['font'] or sp['bbox'][0]<53 or sp['bbox'][2]>559 or sp['bbox'][1]<53 or sp['bbox'][3]>739: raise ValueError('PDF font or margin QA failed')
        d[0].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(preview)
        return max(b[3] for b in d[0].get_text('blocks'))


def prepare(store,id):
    with fs.lock(store.root/("generation-"+str(id)+".lock")):
        return _prepare(store,id)

def _prepare(store,id):
    settings=store.settings();job=store.redact(store.get(id));profile=store.public_profile()
    if not profile['facts']: raise ValueError('Private candidate facts missing')
    selection=select_facts(profile,job);rendered={t['fact_id']:t['rendered'] for t in selection['tailoring']};lookup={f['id']:{**f,'text':rendered[f['id']]} for f in profile['facts']}
    selected=[lookup[i] for i in selection['selected_fact_ids']]
    from personal_jobby.artifacts import begin,publish
    folder,generation,initial=begin(store,id,profile=profile,job=job,settings=settings)
    # One best source bullet from every employment role is mandatory.
    mandatory=[f for f in selected if f['section'] in ('EDUCATION','TECHNICAL SKILLS')]
    roles=set()
    for f in selected:
        if f['section']=='EXPERIENCE' and f['group'] not in roles: mandatory.append(f);roles.add(f['group'])
    chosen=mandatory+[f for f in selected if f not in mandatory and f['section']!='SUMMARY'][:10]
    leading=12.5;body_space=3
    while True:
        content=lines(profile,chosen);pdf(folder/'resume.pdf',content,leading,body_space)
        with fitz.open(folder/'resume.pdf') as d:
            if len(d)==1: break
        removable=[f for f in reversed(chosen) if f not in mandatory]
        if not removable: raise ValueError('Mandatory roles/education/skills exceed one page; manual source-preserving review required')
        chosen.remove(removable[0])
    # Increase line spacing gently when source content is short, without changing margins or fonts.
    bottom=qa(folder/'resume.pdf',folder/'preview.png')
    for trial in (13,13.5,14):
        if bottom>=700: break
        pdf(folder/'resume.pdf',content,trial,body_space)
        with fitz.open(folder/'resume.pdf') as d:
            if len(d)>1:
                pdf(folder/'resume.pdf',content,leading,body_space);break
        leading=trial;bottom=qa(folder/'resume.pdf',folder/'preview.png')
    qa(folder/'resume.pdf',folder/'preview.png');docx(folder/'resume.docx',content,leading,body_space)
    relevant=[f for f in chosen if f['section']=='EXPERIENCE'][:3]
    letter=[('name',profile['name']),('contact',contact_parts(profile)),('body',now()[:10]),('heading',('Re: '+job['title']+' — '+job['company'],'')),('body','Dear Hiring Team,'),('body','Please consider my application for the '+job['title']+' position at '+job['company']+'. My experience relevant to this role includes the following:')]
    letter += [('body','I '+f['text'][0].lower()+f['text'][1:]) for f in relevant]
    letter += [('body','I would welcome the opportunity to discuss how this experience could contribute to your team. Thank you for your time and consideration.'),('body','Sincerely,'),('body',profile['name'])]
    pdf(folder/'cover_letter.pdf',letter,14,6);docx(folder/'cover_letter.docx',letter,14,6);qa(folder/'cover_letter.pdf',folder/'cover_preview.png')
    selection['ranked_fact_ids']=selection['selected_fact_ids'];selection['selected_fact_ids']=[f['id'] for f in chosen]
    selection['gaps']=list(dict.fromkeys(selection['gaps']+job['explanation']))
    selection['evidence']=[{**e,'included':e['fact_id'] in selection['selected_fact_ids'],'text':lookup[e['fact_id']]['text'],'source_excerpt':lookup[e['fact_id']]['excerpt'],'original_text':next(t['original'] for t in selection['tailoring'] if t['fact_id']==e['fact_id'])} for e in selection['evidence']]
    private_write(folder/'evidence.json',json.dumps(selection,indent=2));private_write(folder/'posting.json',json.dumps(store.redact({'original_listing':json.loads(store.events(id)[0]['data']),'current_verified_listing':job}),indent=2))
    meta=dict(created=now(),mode=selection['mode'],selection_mode=selection['selection_mode'],qa_state='pending_visual_review',pdf_pages=1,cover_pdf_pages=1,minimum_font=10,margins_inches=.75,leading=leading,candidate_hash=hashlib.sha256(fs.read_bytes(store.root/'candidate.json')).hexdigest(),source_hashes=profile.get('source_hashes',{}),selected_fact_ids=selection['selected_fact_ids'],docx_qa='Source content equivalent; DOCX pagination requires external visual review')
    meta['artifact_hashes']={n:hashlib.sha256(fs.read_bytes(folder/n)).hexdigest() for n in NAMES if (folder/n).exists() and n!='metadata.json'}
    meta=publish(store,id,folder,generation,meta,initial)
    for p in store.artifact_dir(id).iterdir(): fs.chmod(p)
    with store.connect() as c:
        if job['status'] not in ('applied','interview','offer','rejected','archived'): c.execute('UPDATE jobs SET status=?,updated=? WHERE id=?',('materials_ready',now(),id))
        store.event(c,id,'materials',meta)
    return meta
